"""Excel export of the SQLite brain (attributes/links/calcs — no geometry).

Uses openpyxl when it is already installed; otherwise writes one CSV per
sheet (stdlib). Never installs packages.
"""
from __future__ import annotations

import csv
import json
import sqlite3
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from paths import DB_PATH, QC_ROOT

OUT_BRAIN = QC_ROOT / "quantity-capital-brain.xlsx"
OUT_CLAIMS = QC_ROOT / "quantity-capital-claims.xlsx"


def log(msg: str) -> None:
    print(msg, flush=True)


BRAIN_SHEETS = [
    ("companies",
     ["company_id", "name", "holder", "commodity", "type", "country", "quebec", "ontario", "bc", "claim_count", "sources"],
     "SELECT company_id, name, holder, commodity, company_type, country, quebec_count, ontario_count, bc_count, claim_count, sources_json FROM companies ORDER BY (claim_count IS NULL), claim_count DESC, company_id",
     20000),
    ("tickers", ["ticker", "name", "industry", "market_cap", "px", "cap_asof"],
     "SELECT ticker, name, industry, market_cap, px, cap_asof FROM tickers ORDER BY ticker LIMIT 20000", 20000),
    ("claim_titles",
     ["title_pk", "pack_id", "company_id", "jurisdiction", "role", "claim_id", "holder", "status", "recorded", "expiry", "area_ha", "tenure_type"],
     "SELECT title_pk, pack_id, company_id, jurisdiction, role, claim_id, holder_raw, status, recorded_date, anniversary_or_expiry, area_ha, tenure_type FROM claim_titles LIMIT 20000",
     20000),
    ("claim_company_links", ["pack_id", "company_id", "link_role", "source", "jev_outcome", "jev_score", "jev_confidence"],
     "SELECT pack_id, company_id, link_role, source, jev_outcome, jev_score, jev_confidence FROM claim_company_links", 20000),
    ("trade_size_vs_cap",
     ["trade_id", "ticker", "company_id", "filer", "side", "trade_date", "amount", "mid", "market_cap", "size_bps", "jev_flag"],
     "SELECT trade_id, ticker, company_id, filer, side, trade_date, amount_raw, amount_mid, market_cap, size_bps, jev_flag FROM trade_size_vs_cap WHERE size_bps IS NOT NULL ORDER BY size_bps DESC LIMIT 500",
     20000),
    ("tell_hands", ["filer_id", "name", "chamber", "scored", "hits", "rate", "median", "score", "rank"],
     "SELECT filer_id, name, chamber, scored, hits, rate, median, score, rank FROM tell_hands ORDER BY rank", 20000),
    ("analysis_book", ["ticker", "name", "action", "score", "why", "heat_rank", "list", "jev_ship", "jev_action", "shipped"],
     "SELECT ticker, name, action, score, why, heat_rank, list, jev_ship, jev_action, shipped FROM analysis_candidates WHERE shipped=1 ORDER BY list, score DESC",
     20000),
    ("jev_decisions", ["created_at", "subject_type", "subject_id", "question_id", "primitive", "model", "answer_json"],
     "SELECT created_at, subject_type, subject_id, question_id, primitive, model, substr(answer_json,1,400) FROM jev_decisions ORDER BY id DESC LIMIT 2000",
     20000),
]
CLAIMS_SHEETS = [
    ("claim_titles",
     ["title_pk", "pack_id", "company_id", "jurisdiction", "role", "claim_id", "holder", "status", "recorded", "expiry", "area_ha", "type", "extract"],
     "SELECT title_pk, pack_id, company_id, jurisdiction, role, claim_id, holder_raw, status, recorded_date, anniversary_or_expiry, area_ha, tenure_type, extract_path FROM claim_titles LIMIT 50000",
     50000),
    ("links", ["pack_id", "company_id", "link_role", "source", "jev_outcome", "jev_score"],
     "SELECT pack_id, company_id, link_role, source, jev_outcome, jev_score FROM claim_company_links", 50000),
    ("companies", ["company_id", "name", "holder", "quebec", "ontario", "bc", "claim_count"],
     "SELECT company_id, name, holder, quebec_count, ontario_count, bc_count, claim_count FROM companies WHERE ifnull(claim_count,0)+ifnull(quebec_count,0)+ifnull(ontario_count,0)+ifnull(bc_count,0)>0 ORDER BY claim_count DESC",
     50000),
]


def cell(val):
    if isinstance(val, (dict, list)):
        val = json.dumps(val)
    if isinstance(val, str) and len(val) > 32000:
        val = val[:32000]
    return val


def q(con, sql, args=()):
    return list(con.execute(sql, args))


def write_xlsx(con: sqlite3.Connection, out: Path, sheets: list) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    wb = Workbook()
    for idx, (title, headers, sql, max_rows) in enumerate(sheets):
        ws = wb.active if idx == 0 else wb.create_sheet(title)
        ws.title = title
        for i, h in enumerate(headers, 1):
            ws.cell(1, i, h).font = bold
        for r_i, row in enumerate(q(con, sql)[:max_rows], 2):
            for c_i, val in enumerate(row, 1):
                ws.cell(r_i, c_i, cell(val))
    wb.save(out)
    log(f"Wrote {out}")
    return out


def write_csv_dir(con: sqlite3.Connection, out: Path, sheets: list) -> Path:
    folder = out.with_suffix("")
    folder = folder.with_name(folder.name + "-csv")
    folder.mkdir(parents=True, exist_ok=True)
    for title, headers, sql, max_rows in sheets:
        with (folder / f"{title}.csv").open("w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(headers)
            for row in q(con, sql)[:max_rows]:
                w.writerow([cell(v) for v in row])
    log(f"Wrote {folder}/ ({len(sheets)} CSV files)")
    return folder


def export(con: sqlite3.Connection) -> tuple[Path, Path]:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        log("openpyxl is not installed; writing CSV folders instead (pip install openpyxl for .xlsx).")
        return write_csv_dir(con, OUT_BRAIN, BRAIN_SHEETS), write_csv_dir(con, OUT_CLAIMS, CLAIMS_SHEETS)
    return write_xlsx(con, OUT_BRAIN, BRAIN_SHEETS), write_xlsx(con, OUT_CLAIMS, CLAIMS_SHEETS)


def run(con: sqlite3.Connection) -> None:
    export(con)


def main() -> int:
    if not DB_PATH.exists():
        log(f"missing {DB_PATH}")
        return 1
    con = sqlite3.connect(str(DB_PATH))
    try:
        export(con)
    finally:
        con.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
